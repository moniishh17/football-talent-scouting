"""Position-aware scouting rankings over observed StatsBomb player metrics."""

import argparse
import csv
import math
import re
import time
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from shared.cloud_infra import DistributedStorage, ElasticComputePool


DEFAULT_DATASET = Path(__file__).resolve().parents[1] / "data" / "statsbomb_world_cup_2022_players.csv"
RESULTS_DIR = Path(__file__).resolve().parents[1] / "data" / "results"
DEFAULT_FINAL_REPORT = RESULTS_DIR / "FINAL_TALENT_SCOUTING_RESULTS.xlsx"
DEFAULT_RANKINGS_CSV = DEFAULT_FINAL_REPORT
POSITION_WEIGHTS = {
    "GK": {"saves_per90": 0.35, "clean_sheets_per90": 0.25, "pass_completion": 0.20, "minutes": 0.20},
    "DEF": {"def_actions_per90": 0.35, "pass_completion": 0.20, "goals_assists_per90": 0.15, "minutes": 0.30},
    "MID": {"pass_completion": 0.25, "goals_assists_per90": 0.25, "def_actions_per90": 0.20, "xg_per90": 0.10, "minutes": 0.20},
    "FWD": {"goals_per90": 0.30, "xg_per90": 0.25, "shots_on_target_rate": 0.15, "goals_assists_per90": 0.10, "minutes": 0.20},
}
# Additional FBref metrics are included only when every candidate in that
# position/competition cohort has the field. Core weights retain their ratios.
OPTIONAL_POSITION_WEIGHTS = {
    "GK": {"save_pct": 0.15},
    "DEF": {"xag_per90": 0.05, "progressive_passes_per90": 0.08, "progressive_carries_per90": 0.05},
    "MID": {"xag_per90": 0.10, "progressive_passes_per90": 0.10, "progressive_carries_per90": 0.10},
    "FWD": {"xag_per90": 0.10, "progressive_passes_per90": 0.05, "progressive_carries_per90": 0.10},
}
ALGORITHMS = ("weighted_sum", "topsis", "vikor")
ALGORITHM_LABELS = {"weighted_sum": "WEIGHTED SUM", "topsis": "TOPSIS", "vikor": "VIKOR"}
POSITION_ORDER = ("GK", "DEF", "MID", "FWD")


def create_run_output(label: str = "FINAL_TALENT_SCOUTING_RESULTS", output_dir: str | Path = RESULTS_DIR) -> Path:
    """Return the one stable, regenerated Excel report path."""
    return Path(output_dir) / "FINAL_TALENT_SCOUTING_RESULTS.xlsx"


@dataclass
class PlayerProfile:
    player_id: str
    name: str
    position: str
    metrics: dict[str, float]
    team: str = ""
    country: str = ""
    season: str = ""
    league: str = ""
    source: str = ""


def load_player_pool(path: str | Path = DEFAULT_DATASET, season: str | None = None) -> list[PlayerProfile]:
    """Load a prepared player-statistics CSV; multi-season files require --season."""
    dataset_path = Path(path)
    if not dataset_path.exists():
        raise FileNotFoundError(f"Player statistics CSV not found: {dataset_path}")
    if not dataset_path.is_file():
        raise ValueError(f"Input path is not a file: {dataset_path}")
    with dataset_path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames:
            raise ValueError(f"Input CSV has no header: {dataset_path}")
        headers = {header.strip() for header in reader.fieldnames if header}
        missing = sorted({"player_id", "name", "position", "minutes"} - headers)
        if missing:
            raise ValueError("Missing required columns: " + ", ".join(missing))
        rows = list(reader)
    if "season" in headers:
        seasons = sorted({(row.get("season") or "").strip() for row in rows if (row.get("season") or "").strip()})
        if season is None and len(seasons) > 1:
            raise ValueError("Input has multiple seasons; specify one with --season. Available: " + ", ".join(seasons))
        chosen_season = season or (seasons[0] if seasons else None)
        if chosen_season is not None:
            rows = [row for row in rows if (row.get("season") or "").strip() == chosen_season]
            if not rows:
                raise ValueError(f"Season {chosen_season!r} was not found in {dataset_path}.")
    elif season is not None:
        raise ValueError("--season was supplied, but the input CSV has no season column.")
    if not rows:
        raise ValueError(f"No player rows found in {dataset_path}.")

    players = []
    for line_number, row in enumerate(rows, start=2):
        player_id, name = (row.get("player_id") or "").strip(), (row.get("name") or "").strip()
        if not player_id or not name:
            raise ValueError(f"Row {line_number}: player_id and name must not be blank.")
        position = _normalize_position(row.get("position", ""))
        criteria = set(POSITION_WEIGHTS[position])
        missing_metrics = sorted((criteria - {"goals_assists_per90"}) - headers)
        if "goals_assists_per90" not in headers and not {"goals_per90", "assists_per90"}.issubset(headers):
            missing_metrics.append("goals_assists_per90 (or both goals_per90 and assists_per90)")
        if missing_metrics:
            raise ValueError(f"Row {line_number} ({position}): missing required metrics: " + ", ".join(missing_metrics))

        numeric_fields = (criteria - {"goals_assists_per90"}) | {
            "goals_per90", "assists_per90", "goals_assists_per90", "pass_completion",
            "xg_per90", "shots_on_target_rate", "def_actions_per90", "saves_per90",
            "clean_sheets_per90", "minutes", "save_pct", "xag_per90",
            "progressive_passes_per90", "progressive_carries_per90",
        }
        metrics = {}
        for key in numeric_fields:
            raw = row.get(key)
            if raw in (None, ""):
                if key == "goals_assists_per90" and {"goals_per90", "assists_per90"}.issubset(headers):
                    continue
                if key in criteria:
                    raise ValueError(f"Row {line_number}: required metric {key} is blank.")
                continue
            try:
                value = float(raw)
            except (TypeError, ValueError) as exc:
                raise ValueError(f"Row {line_number}: metric {key} must be numeric; got {raw!r}.") from exc
            if not math.isfinite(value) or value < 0:
                raise ValueError(f"Row {line_number}: metric {key} must be finite and non-negative.")
            metrics[key] = value
        if "goals_assists_per90" not in metrics:
            metrics["goals_assists_per90"] = metrics.get("goals_per90", 0.0) + metrics.get("assists_per90", 0.0)
        for rate in ("pass_completion", "shots_on_target_rate", "save_pct"):
            if rate in metrics and metrics[rate] > 1:
                raise ValueError(f"Row {line_number}: {rate} must be between 0 and 1.")
        players.append(PlayerProfile(player_id, name, position, metrics,
                                     (row.get("team") or "").strip(), (row.get("country") or "").strip(),
                                     (row.get("season") or "").strip(), (row.get("league") or "").strip(),
                                     (row.get("source") or "").strip()))
    seen, duplicate_ids = set(), set()
    for player in players:
        if player.player_id in seen:
            duplicate_ids.add(player.player_id)
        seen.add(player.player_id)
    duplicates = sorted(duplicate_ids)
    if duplicates:
        raise ValueError("Duplicate player_id values after season selection: " + ", ".join(duplicates[:5]))
    return players


def _normalize_position(value: str) -> str:
    """Map common role labels to the four position groups used by the model."""
    label = (value or "").strip().lower()
    # FBref exports may list positions as compact codes or comma-separated roles.
    label = label.split(",", 1)[0].strip()
    if label in {"gk", "goalkeeper", "goalie"} or "goalkeeper" in label:
        return "GK"
    if label in {"df", "def", "defender", "defence", "defense"} or any(
        term in label for term in ("defend", "defence", "defense", "back")
    ):
        return "DEF"
    if label in {"fw", "fwd", "forward", "striker", "winger", "attacker"} or any(
        term in label for term in ("forward", "striker", "wing", "attacker")
    ):
        return "FWD"
    if label in {"mf", "mid", "midfielder", "midfield"} or "mid" in label:
        return "MID"
    raise ValueError(f"Unsupported position {value!r}; use GK, DEF, MID, FWD or a clear full role name.")


def _normalize_group(players: list[PlayerProfile], criteria: tuple[str, ...]) -> list[dict[str, float]]:
    """Min-max normalize within position; ties and constant criteria map to zero."""
    normalized = [dict() for _ in players]
    for criterion in criteria:
        values = [player.metrics.get(criterion, 0.0) for player in players]
        low, high = min(values), max(values)
        for index, value in enumerate(values):
            normalized[index][criterion] = (value - low) / (high - low) if high > low else 0.0
    return normalized


def _effective_weights(players: list[PlayerProfile],
                       base_weights: dict[str, float] | None = None) -> dict[str, float]:
    """Use complete optional metrics for this cohort; never treat missing data as zero."""
    position = players[0].position
    weights = dict(base_weights or POSITION_WEIGHTS[position])
    if base_weights is None:
        for criterion, weight in OPTIONAL_POSITION_WEIGHTS[position].items():
            if all(criterion in player.metrics for player in players):
                weights[criterion] = weight
    total = sum(weights.values())
    return {criterion: weight / total for criterion, weight in weights.items()}


def _weighted_sum(matrix: list[dict[str, float]], weights: dict[str, float]) -> list[float]:
    return [sum(row[c] * weights[c] for c in weights) for row in matrix]


def _topsis(matrix: list[dict[str, float]], weights: dict[str, float]) -> list[float]:
    weighted = [{c: row[c] * weights[c] for c in weights} for row in matrix]
    ideal = dict(weights)
    scores = []
    for row in weighted:
        best_distance = sum((row[c] - ideal[c]) ** 2 for c in weights) ** 0.5
        worst_distance = sum(row[c] ** 2 for c in weights) ** 0.5
        scores.append(worst_distance / (best_distance + worst_distance)
                      if best_distance + worst_distance else 0.0)
    return scores


def _vikor(matrix: list[dict[str, float]], weights: dict[str, float], v: float = 0.5) -> list[float]:
    utility, regret = [], []
    for row in matrix:
        losses = [weights[c] * (1.0 - row[c]) for c in weights]
        utility.append(sum(losses))
        regret.append(max(losses))
    u_min, u_max, r_min, r_max = min(utility), max(utility), min(regret), max(regret)
    scores = []
    for u, r in zip(utility, regret):
        q_u = (u - u_min) / (u_max - u_min) if u_max > u_min else 0.0
        q_r = (r - r_min) / (r_max - r_min) if r_max > r_min else 0.0
        scores.append(1.0 - (v * q_u + (1.0 - v) * q_r))
    return scores


def _rank_group(players: list[PlayerProfile], weights_override: dict[str, float] | None = None) -> dict[str, list[dict]]:
    position = players[0].position
    weights = _effective_weights(players, weights_override)
    matrix = _normalize_group(players, tuple(weights))
    method_scores = {
        "weighted_sum": _weighted_sum(matrix, weights),
        "topsis": _topsis(matrix, weights),
        "vikor": _vikor(matrix, weights),
    }
    result = {}
    for algorithm, scores in method_scores.items():
        ordered = sorted(zip(players, scores), key=lambda item: (-item[1], item[0].player_id))
        result[algorithm] = [
            {"rank": rank, "player_id": player.player_id, "name": player.name,
             "team": player.team, "country": player.country, "season": player.season,
             "league": player.league, "source": player.source, "position": position,
             "score": round(score, 6), "metrics": player.metrics,
             "criteria_weights": weights}
            for rank, (player, score) in enumerate(ordered, 1)
        ]
    return result


def _spearman(left: list[dict], right: list[dict]) -> float:
    right_rank = {entry["player_id"]: entry["rank"] for entry in right}
    n = len(left)
    if n < 2:
        return 1.0
    squared_diff = sum((entry["rank"] - right_rank[entry["player_id"]]) ** 2 for entry in left)
    return 1.0 - (6.0 * squared_diff) / (n * (n * n - 1))


def run_hpc_processing(players: list[PlayerProfile], compute: ElasticComputePool) -> dict:
    """Rank each position group with WSM, TOPSIS, and VIKOR in parallel."""
    grouped: dict[str, list[PlayerProfile]] = {}
    for player in players:
        if player.position not in POSITION_WEIGHTS:
            continue
        grouped.setdefault(player.position, []).append(player)
    compute.scale(pending_jobs=len(grouped))
    futures = {position: compute.submit(_rank_group, group) for position, group in grouped.items()}
    return {position: future.result() for position, future in futures.items()}


def run_output_intelligence(rankings: dict, players: list[PlayerProfile], storage: DistributedStorage,
                            dataset_path: str | Path = DEFAULT_DATASET,
                            rankings_output: str | Path | None = None,
                            min_minutes: float = 180) -> dict:
    top = {
        position: {algorithm: ranked[:5] for algorithm, ranked in methods.items()}
        for position, methods in rankings.items()
    }
    agreement = {
        position: {
            f"{left}_vs_{right}_spearman": round(_spearman(methods[left], methods[right]), 4)
            for index, left in enumerate(ALGORITHMS)
            for right in ALGORITHMS[index + 1:]
        }
        for position, methods in rankings.items()
    }
    for position, methods in rankings.items():
        for algorithm, ranked in methods.items():
            for entry in ranked:
                storage.put(f"scouting:{algorithm}:{entry['player_id']}", entry)
    default_source = Path(dataset_path).resolve() == DEFAULT_DATASET.resolve()
    summary = {
        "data_source": ("StatsBomb Open Data, FIFA World Cup 2022 (competition 43, season 106)"
                        if default_source else (players[0].source or f"Prepared player-statistics CSV: {Path(dataset_path).name}")),
        "total_players_scouted": len(players),
        "ranking_scope": f"within_primary_position; minimum {min_minutes:g} input minutes",
        "algorithms": list(ALGORITHMS),
        "position_weights": {
            position: methods[ALGORITHMS[0]][0]["criteria_weights"]
            for position, methods in rankings.items() if methods.get(ALGORITHMS[0])
        },
        "normalization": "min-max within position group; higher is better for all criteria",
        "full_rankings_file": str(rankings_output),
        "rankings_workbook_file": str(rankings_output),
        "top_candidates_by_algorithm_and_position": top,
        "rankings": rankings,
        "rank_agreement_spearman": agreement,
        "research_note": (
            "This tournament dataset has performance observations but no independent future-scouting "
            "outcome label. Agreement is descriptive and cannot establish predictive superiority. "
            "The weights are researcher-defined starting values, not elicited or learned."
        ),
    }
    storage.put("scouting:summary", summary)
    return summary


def run_scouting_pipeline(compute: ElasticComputePool, storage: DistributedStorage,
                          dataset_path: str | Path = DEFAULT_DATASET,
                          rankings_output: str | Path | None = None,
                          season: str | None = None,
                          min_minutes: float = 180,
                          output_dir: str | Path | None = None) -> dict:
    """Load real player metrics and publish position-specific comparison rankings."""
    if not math.isfinite(min_minutes) or min_minutes < 0:
        raise ValueError("Minimum minutes must be a finite, non-negative number.")
    players = load_player_pool(dataset_path, season=season)
    players = [player for player in players if player.metrics["minutes"] >= min_minutes]
    if not players:
        raise ValueError(f"No players meet the minimum of {min_minutes:g} minutes in {dataset_path}.")
    rankings_output = Path(rankings_output) if rankings_output else create_run_output(output_dir=output_dir or RESULTS_DIR)
    # The named deliverable is always an Excel workbook; CSVs are no longer
    # emitted as parallel result files.
    workbook_path = rankings_output if rankings_output.suffix.lower() == ".xlsx" else rankings_output.with_suffix(".xlsx")
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    storage.put("scouting:raw_ingest_count", len(players))
    started = time.perf_counter()
    rankings = run_hpc_processing(players, compute)
    from processing.temporal_validation import DATA_PATH as SELECTION_DATA_PATH, evaluate, weight_sensitivity
    from processing.algorithm_selection_analysis import compare as compare_algorithms

    selection_validation = evaluate(SELECTION_DATA_PATH)
    selection_sensitivity = weight_sensitivity(SELECTION_DATA_PATH)
    selection_rows = compare_algorithms(
        SELECTION_DATA_PATH, None,
        validation_rows=selection_validation, sensitivity_rows=selection_sensitivity,
        print_results=False,
    )
    selection_summaries = [row for row in selection_rows if row["record_type"] == "algorithm_summary"]
    winner = min(selection_summaries, key=lambda row: int(row["primary_rank"]))
    winner_key = next(key for key, label in ALGORITHM_LABELS.items()
                      if label.casefold() == winner["algorithm"].casefold())
    summary = run_output_intelligence(rankings, players, storage, dataset_path, workbook_path, min_minutes)
    summary["processing_seconds"] = round(time.perf_counter() - started, 4)
    paired_differences = [row for row in selection_rows
                          if row["record_type"] == "paired_difference"
                          and row["algorithm"] == winner["algorithm"]]
    uncertain_comparisons = [row for row in paired_differences
                             if float(row["paired_bootstrap_95ci_low"]) <= 0
                             <= float(row["paired_bootstrap_95ci_high"])]
    summary["selected_algorithm"] = winner_key
    summary["selected_algorithm_label"] = winner["algorithm"]
    summary["selected_rankings_file"] = str(workbook_path)
    summary["validation_rows"] = selection_validation
    summary["sensitivity_rows"] = selection_sensitivity
    summary["selection_rows"] = selection_rows
    summary["algorithm_selection"] = {
        "method": winner["algorithm"],
        "primary_metric": "future performance proxy NDCG@10",
        "mean_score": float(winner["primary_mean"]),
        "mean_spearman": float(winner["future_performance_spearman_mean"]),
        "validation_folds": int(winner["performance_folds"]),
        "selection_dataset": f"La Liga player-seasons; {SELECTION_DATA_PATH.name}",
        "paired_comparisons": [{
            "against": row["comparison_against"],
            "mean_difference": float(row["paired_primary_mean_difference"]),
            "ci95_low": float(row["paired_bootstrap_95ci_low"]),
            "ci95_high": float(row["paired_bootstrap_95ci_high"]),
        } for row in paired_differences],
        "interpretation": (
            "Provisional selection by the pre-specified mean NDCG@10 rule. The paired 95% interval "
            "against at least one other method includes zero, so superiority is not established."
            if uncertain_comparisons else
            "Selected by the pre-specified mean NDCG@10 rule; paired fold intervals exclude zero "
            "against the other methods in this pilot."
        ),
    }
    _save_rankings_workbook(rankings, summary, players, workbook_path)
    summary["previous_reports_archived"] = _archive_previous_reports(workbook_path)
    storage.put("scouting:summary", summary)
    return summary


def _save_full_rankings(rankings: dict, output_path: str | Path,
                        algorithms: tuple[str, ...] = ALGORITHMS) -> None:
    """Export every candidate's score and component metrics for review."""
    metric_fields = (
        "minutes", "pass_completion", "goals_per90", "assists_per90", "xg_per90",
        "shots_on_target_rate", "def_actions_per90", "saves_per90", "clean_sheets_per90",
        "goals_assists_per90", "save_pct", "xag_per90", "progressive_passes_per90",
        "progressive_carries_per90",
    )
    fields = ("algorithm", "position", "rank", "player_id", "name", "team", "country",
              "season", "league", "source", "score", *metric_fields)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for algorithm in algorithms:
            for position in POSITION_ORDER:
                if position not in rankings:
                    continue
                entries = rankings[position].get(algorithm, [])
                for entry in entries:
                    row = {key: entry[key] for key in fields if key in entry}
                    row.update(entry["metrics"])
                    row["algorithm"] = algorithm
                    writer.writerow(row)


def _save_rankings_workbook(rankings: dict, summary: dict, players: list[PlayerProfile],
                            output_path: str | Path) -> None:
    """Write a reader-friendly workbook with one full ranking sheet per method."""
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    overview.sheet_view.showGridLines = False
    workbook.properties.title = "Talent Scouting Ranking Results"
    workbook.properties.subject = "Position-aware comparison of Weighted Sum, TOPSIS and VIKOR"
    workbook.properties.creator = "Talent Scouting Research Project"

    navy, teal, pale, ink, gray = "17324D", "138A8A", "EAF3F7", "243447", "667788"
    white = "FFFFFF"
    header_fill = PatternFill("solid", fgColor=navy)
    title_fill = PatternFill("solid", fgColor=navy)
    accent_fill = PatternFill("solid", fgColor=teal)
    thin_gray = Side(style="thin", color="D8E1E8")

    def style_title(sheet, title, subtitle, end_col=8):
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
        title_cell = sheet.cell(1, 1, title)
        title_cell.font = Font(name="Aptos Display", size=19, bold=True, color=white)
        title_cell.fill = title_fill
        title_cell.alignment = Alignment(vertical="center")
        for column in range(1, end_col + 1):
            sheet.cell(1, column).fill = title_fill
        sheet.row_dimensions[1].height = 34
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
        subtitle_cell = sheet.cell(2, 1, subtitle)
        subtitle_cell.font = Font(name="Aptos", size=10, color=gray, italic=True)
        subtitle_cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[2].height = 30
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 90

    def style_header(sheet, row, start_col, end_col):
        for col in range(start_col, end_col + 1):
            cell = sheet.cell(row, col)
            cell.fill = header_fill
            cell.font = Font(name="Aptos", bold=True, color=white, size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="medium", color=teal))
        sheet.row_dimensions[row].height = 30

    style_title(overview, "TALENT SCOUTING | RANKING REPORT",
                "Three position-aware multi-criteria methods applied to the same historical player cohort", 8)
    overview.column_dimensions["A"].width = 26
    overview.column_dimensions["B"].width = 36
    for col in "CDEFGH":
        overview.column_dimensions[col].width = 18
    overview["A4"] = "STUDY SCOPE"
    overview["A4"].fill = accent_fill
    overview["A4"].font = Font(bold=True, color=white)
    scope = [
        ("Source / cohort", summary["data_source"]),
        ("Players included", summary["total_players_scouted"]),
        ("Eligibility threshold", summary["ranking_scope"]),
        ("Normalization", summary["normalization"]),
        ("Methods", "Weighted Sum; TOPSIS; VIKOR"),
        ("Output structure", "One workbook contains all three full rankings and the validation evidence."),
    ]
    for row_index, (label, value) in enumerate(scope, start=5):
        overview.cell(row_index, 1, label).font = Font(bold=True, color=ink)
        overview.cell(row_index, 2, value).alignment = Alignment(wrap_text=True, vertical="top")
        overview.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=8)
        overview.cell(row_index, 2).fill = PatternFill("solid", fgColor=pale if row_index % 2 else white)
        overview.row_dimensions[row_index].height = 27 if row_index != 10 else 34

    selection = summary["algorithm_selection"]
    overview["A12"] = "CORE RESULT | PROVISIONAL BEST METHOD"
    overview["A12"].fill = accent_fill
    overview["A12"].font = Font(bold=True, color=white)
    overview.merge_cells("B12:H12")
    overview["B12"] = f"{selection['method']} — selected by mean {selection['primary_metric']}"
    overview["B12"].font = Font(bold=True, size=13, color=navy)
    overview["B12"].fill = PatternFill("solid", fgColor="DDEFEF")
    selection_rows = [
        ("Mean validation score / folds", f"{selection['mean_score']:.4f} / {selection['validation_folds']}"),
        ("Selection cohort", selection["selection_dataset"]),
        ("Evidence note", selection["interpretation"]),
        ("Highlighted method sheet", f"{selection['method']} (gold tab)"),
    ]
    for row_index, (label, value) in enumerate(selection_rows, start=13):
        overview.cell(row_index, 1, label).font = Font(bold=True, color=ink)
        overview.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=8)
        overview.cell(row_index, 2, value).alignment = Alignment(wrap_text=True, vertical="top")
        overview.cell(row_index, 2).fill = PatternFill("solid", fgColor=pale if row_index % 2 else white)
        overview.row_dimensions[row_index].height = 32 if row_index in (15, 16) else 24

    counts = {position: sum(player.position == position for player in players) for position in POSITION_ORDER}
    overview["A18"] = "COHORT COVERAGE"
    overview["A18"].fill = accent_fill
    overview["A18"].font = Font(bold=True, color=white)
    coverage_headers = ["Position", "Candidates", "Weighted Sum", "TOPSIS", "VIKOR"]
    for col, value in enumerate(coverage_headers, start=1):
        overview.cell(19, col, value)
    style_header(overview, 19, 1, len(coverage_headers))
    for row_index, position in enumerate(POSITION_ORDER, start=20):
        if position not in counts or not counts[position]:
            continue
        values = [position, counts[position], counts[position], counts[position], counts[position]]
        for col, value in enumerate(values, start=1):
            cell = overview.cell(row_index, col, value)
            cell.fill = PatternFill("solid", fgColor=pale if row_index % 2 == 0 else white)
            cell.border = Border(bottom=thin_gray)

    agreement_start = 25
    overview.cell(agreement_start, 1, "RANK AGREEMENT BY POSITION (SPEARMAN ρ)")
    overview.cell(agreement_start, 1).fill = accent_fill
    overview.cell(agreement_start, 1).font = Font(bold=True, color=white)
    agreement_headers = ["Position", "Weighted Sum vs TOPSIS", "Weighted Sum vs VIKOR", "TOPSIS vs VIKOR"]
    for col, value in enumerate(agreement_headers, start=1):
        overview.cell(agreement_start + 1, col, value)
    style_header(overview, agreement_start + 1, 1, len(agreement_headers))
    for row_index, position in enumerate(POSITION_ORDER, start=agreement_start + 2):
        pairs = summary["rank_agreement_spearman"].get(position, {})
        if not pairs:
            continue
        values = [position, pairs.get("weighted_sum_vs_topsis_spearman"),
                  pairs.get("weighted_sum_vs_vikor_spearman"), pairs.get("topsis_vs_vikor_spearman")]
        for col, value in enumerate(values, start=1):
            cell = overview.cell(row_index, col, value)
            if col > 1 and value is not None:
                cell.number_format = "0.0000"
            cell.fill = PatternFill("solid", fgColor=pale if row_index % 2 == 0 else white)
            cell.border = Border(bottom=thin_gray)

    note_row = 33
    overview.cell(note_row, 1, "INTERPRETATION")
    overview.cell(note_row, 1).fill = accent_fill
    overview.cell(note_row, 1).font = Font(bold=True, color=white)
    overview.merge_cells(start_row=note_row + 1, start_column=1, end_row=note_row + 3, end_column=8)
    overview.cell(note_row + 1, 1, summary["research_note"] +
                  " Rank agreement measures similarity between methods; it does not identify a best method. "
                  "Use the temporal-validation and sensitivity analyses as separate evidence, and state their limits.")
    overview.cell(note_row + 1, 1).alignment = Alignment(wrap_text=True, vertical="top")
    overview.cell(note_row + 1, 1).fill = PatternFill("solid", fgColor=pale)
    overview.freeze_panes = "A4"
    overview.sheet_properties.pageSetUpPr.fitToPage = True
    overview.page_setup.orientation = "landscape"
    overview.page_setup.fitToWidth = 1
    overview.page_setup.fitToHeight = 1
    overview.print_area = "A1:H36"

    weights_sheet = workbook.create_sheet("Weights")
    style_title(weights_sheet, "MODEL WEIGHTS & SCORING NOTES",
                "Position-specific normalized weights actually used for this cohort", 5)
    weight_headers = ["Position", "Criterion", "Normalized weight", "Metric direction", "Scoring note"]
    for col, value in enumerate(weight_headers, start=1):
        weights_sheet.cell(4, col, value)
    style_header(weights_sheet, 4, 1, len(weight_headers))
    weight_row = 5
    for position in POSITION_ORDER:
        for criterion, weight in summary.get("position_weights", {}).get(position, {}).items():
            row_values = [position, criterion, weight, "Higher is better", "Min-max normalized within position"]
            for col, value in enumerate(row_values, start=1):
                cell = weights_sheet.cell(weight_row, col, value)
                cell.border = Border(bottom=thin_gray)
                if weight_row % 2:
                    cell.fill = PatternFill("solid", fgColor=pale)
            weights_sheet.cell(weight_row, 3).number_format = "0.0000"
            weight_row += 1
    weights_sheet.freeze_panes = "A5"
    for col, width in {"A": 14, "B": 30, "C": 22, "D": 20, "E": 42}.items():
        weights_sheet.column_dimensions[col].width = width
    if weight_row > 5:
        table = Table(displayName="ModelWeights", ref=f"A4:E{weight_row - 1}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True,
                                              showColumnStripes=False, showFirstColumn=False, showLastColumn=False)
        weights_sheet.add_table(table)

    metric_fields = ("minutes", "pass_completion", "goals_per90", "assists_per90", "xg_per90",
                     "shots_on_target_rate", "def_actions_per90", "saves_per90", "clean_sheets_per90",
                     "goals_assists_per90", "save_pct", "xag_per90", "progressive_passes_per90",
                     "progressive_carries_per90")
    method_colors = {"weighted_sum": "138A8A", "topsis": "3975A8", "vikor": "8556A1"}
    display_names = {"weighted_sum": "Weighted Sum", "topsis": "TOPSIS", "vikor": "VIKOR"}
    for algorithm in ALGORITHMS:
        sheet = workbook.create_sheet(display_names[algorithm])
        is_selected = algorithm == summary["selected_algorithm"]
        sheet.sheet_properties.tabColor = "D0A52B" if is_selected else method_colors[algorithm]
        headers = ["Position", "Rank", "Player ID", "Player", "Team", "Country", "Season", "Competition",
                   "Score", *metric_fields]
        title_suffix = " — SELECTED BEST METHOD" if is_selected else ""
        style_title(sheet, f"FULL RANKING | {display_names[algorithm].upper()}{title_suffix}",
                    "Ranked separately within each position; score values are method-specific and should not be compared across methods.",
                    len(headers))
        header_row = 4
        for col, value in enumerate(headers, start=1):
            sheet.cell(header_row, col, value.replace("_", " ").title())
        style_header(sheet, header_row, 1, len(headers))
        data_row = header_row + 1
        for position in POSITION_ORDER:
            for entry in rankings.get(position, {}).get(algorithm, []):
                values = [entry["position"], entry["rank"], entry["player_id"], entry["name"], entry["team"],
                          entry["country"], entry["season"], entry["league"], entry["score"]]
                values.extend(entry["metrics"].get(field) for field in metric_fields)
                for col, value in enumerate(values, start=1):
                    cell = sheet.cell(data_row, col, value)
                    cell.font = Font(name="Aptos", size=10, color=ink)
                    if col in (2, 9) or col >= 10:
                        cell.number_format = "0.0000" if col > 2 else "0"
                    if col == 9:
                        cell.font = Font(name="Aptos", size=10, bold=True, color=method_colors[algorithm])
                data_row += 1
        if data_row > header_row + 1:
            table = Table(displayName=f"{algorithm.title().replace('_', '')}Rankings",
                          ref=f"A{header_row}:{openpyxl_column_letter(len(headers))}{data_row - 1}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True,
                                                  showColumnStripes=False, showFirstColumn=False, showLastColumn=False)
            sheet.add_table(table)
            sheet.conditional_formatting.add(f"I{header_row + 1}:I{data_row - 1}",
                ColorScaleRule(start_type="min", start_color="F4CCCC", mid_type="percentile", mid_value=50,
                               mid_color="FFF2CC", end_type="max", end_color="D9EAD3"))
        sheet.freeze_panes = "D5"
        sheet.auto_filter.ref = f"A{header_row}:{openpyxl_column_letter(len(headers))}{data_row - 1}"
        widths = {"A": 13, "B": 9, "C": 16, "D": 28, "E": 22, "F": 18, "G": 15, "H": 25, "I": 13}
        for col, width in widths.items():
            sheet.column_dimensions[col].width = width
        for col in range(10, len(headers) + 1):
            sheet.column_dimensions[openpyxl_column_letter(col)].width = 19
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:4"

    def add_evidence_sheet(title, rows, table_name):
        if not rows:
            return
        sheet = workbook.create_sheet(title)
        headers = list(rows[0])
        style_title(sheet, title.upper().replace("_", " "),
                    "Supporting research evidence used in the algorithm comparison; row-level results.",
                    max(3, min(len(headers), 10)))
        for col, header in enumerate(headers, start=1):
            sheet.cell(4, col, header.replace("_", " ").title())
        style_header(sheet, 4, 1, len(headers))
        for row_number, record in enumerate(rows, start=5):
            for col, header in enumerate(headers, start=1):
                value = record.get(header)
                try:
                    if value not in (None, ""):
                        value = float(value) if "." in str(value) else int(value)
                except (TypeError, ValueError):
                    pass
                cell = sheet.cell(row_number, col, value)
                cell.font = Font(name="Aptos", size=10, color=ink)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_number % 2:
                    cell.fill = PatternFill("solid", fgColor=pale)
        last_row = 4 + len(rows)
        last_col = openpyxl_column_letter(len(headers))
        table = Table(displayName=table_name, ref=f"A4:{last_col}{last_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True,
                                              showColumnStripes=False, showFirstColumn=False, showLastColumn=False)
        sheet.add_table(table)
        sheet.freeze_panes = "A5"
        sheet.auto_filter.ref = f"A4:{last_col}{last_row}"
        for col, header in enumerate(headers, start=1):
            sheet.column_dimensions[openpyxl_column_letter(col)].width = min(34, max(16, len(header) + 3))

    add_evidence_sheet("Algorithm Comparison", summary.get("selection_rows", []), "AlgorithmComparison")
    add_evidence_sheet("Temporal Validation", summary.get("validation_rows", []), "TemporalValidation")
    add_evidence_sheet("Weight Sensitivity", summary.get("sensitivity_rows", []), "WeightSensitivity")

    workbook.active = 0
    workbook.save(output)


def openpyxl_column_letter(number: int) -> str:
    """Convert a 1-based column index to its Excel letter representation."""
    letters = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _archive_previous_reports(current_xlsx: Path) -> int:
    """Move prior exports and scratch CSVs out of the active results folder."""
    results_dir = current_xlsx.parent
    data_dir = results_dir.parent
    archive_dir = data_dir / "archive" / "previous_rankings"
    current = current_xlsx.resolve()
    candidates = list(results_dir.iterdir()) if results_dir.exists() else []
    moved = 0
    for path in candidates:
        if not path.is_file() or path.resolve() == current or path.name.startswith("~$"):
            continue
        archive_dir.mkdir(parents=True, exist_ok=True)
        destination = archive_dir / path.name
        suffix = 2
        while destination.exists():
            destination = archive_dir / f"{path.stem}_{suffix}{path.suffix}"
            suffix += 1
        try:
            path.replace(destination)
        except OSError:
            # Leave a file that another program currently holds open; it is not
            # read by this pipeline and can be archived on a later successful run.
            continue
        moved += 1
    return moved


def format_rankings_for_console(rankings: dict) -> str:
    """Render in-memory results as three complete, clearly separated blocks."""
    if not rankings:
        return "No ranking rows were written."
    sections = []
    for algorithm_index, algorithm in enumerate(ALGORITHMS, start=1):
        if not any(methods.get(algorithm) for methods in rankings.values()):
            continue
        lines = ["=" * 88, f"ALGORITHM {algorithm_index} OF {len(ALGORITHMS)} | {ALGORITHM_LABELS[algorithm]}"]
        for position in POSITION_ORDER:
            position_rows = rankings.get(position, {}).get(algorithm, [])
            if not position_rows:
                continue
            lines.extend(("", f"{position} rankings ({len(position_rows)} players)"))
            for row in position_rows:
                metadata = " | ".join(value for value in (row.get("season", ""), row.get("league", "")) if value)
                metrics = ", ".join(f"{key}={value:.3f}" for key, value in row["metrics"].items())
                context = f" [{row['team']}]" if row.get("team") else ""
                suffix = f" | {metadata}" if metadata else ""
                lines.append(
                    f"  {int(row['rank']):>3}. {row['name']}{context} | score={float(row['score']):.6f}"
                    f"{suffix} | {metrics}"
                )
        sections.append("\n".join(lines))
    return "\n\n\n".join(sections)


def _run_cli():
    """Run the scouting ranking pipeline on a user-selected prepared CSV."""
    parser = argparse.ArgumentParser(description="Rank historical player statistics with all three MCDA methods.")
    parser.add_argument("--input", default=str(DEFAULT_DATASET), help="prepared player-statistics CSV")
    parser.add_argument("--output", help=f"final Excel workbook path (default: {DEFAULT_FINAL_REPORT})")
    parser.add_argument("--season", help="select one season when the input CSV contains several seasons")
    parser.add_argument("--min-minutes", type=float, default=180,
                        help="minimum minutes needed for ranking (default: 180)")
    args = parser.parse_args()

    storage = DistributedStorage()
    compute = ElasticComputePool(min_workers=2, max_workers=8)
    try:
        summary = run_scouting_pipeline(
            compute, storage, args.input, args.output, args.season, args.min_minutes,
        )
        print(f"Input: {args.input}")
        if args.season:
            print(f"Season: {args.season}")
        print(f"Ranked players: {summary['total_players_scouted']}")
        print(f"Methods: {', '.join(summary['algorithms'])}")
        output = summary["rankings_workbook_file"]
        print(f"Final consolidated report: {output}")
        print("\n" + format_rankings_for_console(summary["rankings"]))
    finally:
        compute.shutdown()


if __name__ == "__main__":
    _run_cli()
