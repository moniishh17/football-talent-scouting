import csv
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from processing.prepare_statsbomb_dataset import (
    BUNDLED_DATASET_PATH, _clock_minutes, _position_group, prepare_dataset,
)
from processing.scouting_channel import (
    ALGORITHMS,
    DEFAULT_DATASET,
    PlayerProfile,
    _rank_group,
    create_run_output,
    format_rankings_for_console,
    load_player_pool,
    run_hpc_processing,
    run_scouting_pipeline,
)
from processing.fbref_import import prepare_fbref_exports
from shared.cloud_infra import DistributedStorage, ElasticComputePool


class DatasetPreparationHelpersTest(unittest.TestCase):
    def test_dataset_preparation_writes_to_requested_path(self):
        prepared_row = {"player_id":"1", "name":"Player", "position":"FWD", "minutes":220}
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parents[1]) as directory:
            output = Path(directory) / "generated" / "players.csv"
            with patch("processing.prepare_statsbomb_dataset.aggregate_world_cup", return_value=[prepared_row]):
                count = prepare_dataset(output)
            self.assertEqual(count, 1)
            with output.open(newline="", encoding="utf-8") as handle:
                self.assertEqual(list(csv.DictReader(handle)), [{key: str(value) for key, value in prepared_row.items()}])

    def test_dataset_preparation_refuses_to_replace_bundled_input(self):
        with self.assertRaisesRegex(ValueError, "Refusing to overwrite"):
            prepare_dataset(BUNDLED_DATASET_PATH)

    def test_position_mapping(self):
        self.assertEqual(_position_group("Goalkeeper"), "GK")
        self.assertEqual(_position_group("Left Center Back"), "DEF")
        self.assertEqual(_position_group("Center Midfield"), "MID")
        self.assertEqual(_position_group("Right Wing"), "FWD")

    def test_clock_parser(self):
        self.assertAlmostEqual(_clock_minutes("64:30"), 64.5)
        self.assertEqual(_clock_minutes(None, 95), 95)


class ScoutingRankingTest(unittest.TestCase):
    def test_fbref_exports_prepare_position_metrics_and_keep_source_metadata(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parents[1]) as directory:
            root = Path(directory)
            standard = root / "standard.csv"
            shooting = root / "shooting.csv"
            passing = root / "passing.csv"
            defense = root / "defense.csv"
            output = root / "prepared.csv"
            tables = {
                standard: (["Player", "Nation", "Pos", "Squad", "Comp", "Min", "Gls", "Ast"],
                           ["Research Player", "eng ENG", "MF", "Test FC", "Test League", "900", "5", "4"]),
                shooting: (["Player", "Squad", "Comp", "xG", "Sh", "SoT"],
                           ["Research Player", "Test FC", "Test League", "6", "20", "8"]),
                passing: (["Player", "Squad", "Comp", "Cmp", "Att", "xAG", "PrgP", "PrgC"],
                          ["Research Player", "Test FC", "Test League", "500", "600", "3", "40", "30"]),
                defense: (["Player", "Squad", "Comp", "Tkl", "Int"],
                          ["Research Player", "Test FC", "Test League", "10", "5"]),
            }
            for path, (fields, row) in tables.items():
                with path.open("w", newline="", encoding="utf-8") as handle:
                    writer = csv.writer(handle)
                    writer.writerow(fields)
                    writer.writerow(row)

            self.assertEqual(prepare_fbref_exports(list(tables), output, "2022-2023"), 1)
            player = load_player_pool(output, season="2022-2023")[0]
            self.assertAlmostEqual(player.metrics["goals_per90"], 0.5)
            self.assertAlmostEqual(player.metrics["assists_per90"], 0.4)
            self.assertAlmostEqual(player.metrics["xg_per90"], 0.6)
            self.assertAlmostEqual(player.metrics["pass_completion"], 5 / 6)
            self.assertAlmostEqual(player.metrics["def_actions_per90"], 1.5)
            self.assertAlmostEqual(player.metrics["xag_per90"], 0.3)
            self.assertAlmostEqual(player.metrics["progressive_passes_per90"], 4.0)
            self.assertAlmostEqual(player.metrics["progressive_carries_per90"], 3.0)
            self.assertEqual(player.league, "Test League")
            self.assertEqual(player.source, "FBref user-exported CSV")
            compute = ElasticComputePool(min_workers=1, max_workers=2)
            rankings_path = root / "rankings.csv"
            try:
                summary = run_scouting_pipeline(compute, DistributedStorage(), output, rankings_path,
                                                 "2022-2023", 180)
            finally:
                compute.shutdown()
            ranking_rows = [entry for methods in summary["rankings"].values()
                            for algorithm in ALGORITHMS for entry in methods[algorithm]]
            self.assertEqual({algorithm for methods in summary["rankings"].values() for algorithm in methods}, set(ALGORITHMS))
            self.assertTrue(all(row["league"] == "Test League" for row in ranking_rows))
            self.assertTrue(all(row["season"] == "2022-2023" for row in ranking_rows))
            weights = summary["position_weights"]["MID"]
            self.assertIn("xag_per90", weights)
            self.assertIn("progressive_passes_per90", weights)
            self.assertAlmostEqual(sum(weights.values()), 1.0)
            self.assertIn(summary["selected_algorithm"], ALGORITHMS)
            self.assertEqual(Path(summary["selected_rankings_file"]), Path(summary["rankings_workbook_file"]))
            self.assertTrue(Path(summary["rankings_workbook_file"]).exists())
            for algorithm in ALGORITHMS:
                method_rows = [entry for methods in summary["rankings"].values()
                               for entry in methods[algorithm]]
                self.assertEqual(len(method_rows), len({row["player_id"] for row in method_rows}))
            book = load_workbook(summary["rankings_workbook_file"], read_only=False)
            self.assertEqual(book.sheetnames, ["Overview", "Weights", "Weighted Sum", "TOPSIS", "VIKOR",
                                               "Algorithm Comparison", "Temporal Validation", "Weight Sensitivity"])
            self.assertEqual(book["Weighted Sum"].max_row - 4, 1)
            self.assertEqual(book["TOPSIS"].max_row - 4, 1)
            self.assertEqual(book["VIKOR"].max_row - 4, 1)
            self.assertIn("SELECTED BEST METHOD", book[summary["selected_algorithm_label"]].cell(1, 1).value)
            self.assertGreater(book["Temporal Validation"].max_row, 4)
            book.close()
            formatted = format_rankings_for_console(summary["rankings"])
            self.assertLess(formatted.index("ALGORITHM 1 OF 3"), formatted.index("ALGORITHM 2 OF 3"))
            self.assertLess(formatted.index("ALGORITHM 2 OF 3"), formatted.index("ALGORITHM 3 OF 3"))

    def test_default_rankings_path_is_stable_for_fresh_replacement(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parents[1]) as directory:
            first = create_run_output("scouting_rankings", directory)
            second = create_run_output("scouting_rankings", directory)
        self.assertEqual(first, second)
        self.assertEqual(first.name, "FINAL_TALENT_SCOUTING_RESULTS.xlsx")

    def test_historical_csv_can_select_a_season_and_export_all_algorithms(self):
        fields = ["season", "player_id", "name", "position", "team", "country", "minutes",
                  "goals_per90", "assists_per90", "xg_per90", "shots_on_target_rate"]
        rows = [
            {"season":"2019/2020", "player_id":"p1", "name":"Forward One", "position":"Forward",
             "team":"Club A", "country":"A", "minutes":"400", "goals_per90":"0.8",
             "assists_per90":"0.2", "xg_per90":"0.7", "shots_on_target_rate":"0.5"},
            {"season":"2019/2020", "player_id":"p2", "name":"Forward Two", "position":"FWD",
             "team":"Club B", "country":"B", "minutes":"350", "goals_per90":"0.3",
             "assists_per90":"0.1", "xg_per90":"0.4", "shots_on_target_rate":"0.4"},
            {"season":"2020/2021", "player_id":"p1", "name":"Forward One", "position":"Forward",
             "team":"Club A", "country":"A", "minutes":"410", "goals_per90":"0.7",
             "assists_per90":"0.2", "xg_per90":"0.6", "shots_on_target_rate":"0.5"},
        ]
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parents[1]) as directory:
            source = Path(directory) / "player_seasons.csv"
            output = Path(directory) / "rankings.csv"
            output.write_text("stale rankings should be replaced", encoding="utf-8")
            output.with_suffix(".xlsx").write_bytes(b"stale workbook should be replaced")
            with source.open("w", newline="", encoding="utf-8") as handle:
                writer = csv.DictWriter(handle, fieldnames=fields)
                writer.writeheader()
                writer.writerows(rows)

            with self.assertRaisesRegex(ValueError, "multiple seasons"):
                load_player_pool(source)
            players = load_player_pool(source, season="2019/2020")
            self.assertEqual([player.position for player in players], ["FWD", "FWD"])
            self.assertAlmostEqual(players[0].metrics["goals_assists_per90"], 1.0)

            compute = ElasticComputePool(min_workers=1, max_workers=2)
            try:
                summary = run_scouting_pipeline(
                    compute, DistributedStorage(), source, output, "2019/2020", 180,
                )
            finally:
                compute.shutdown()
            ranked_book = load_workbook(summary["rankings_workbook_file"], read_only=True)
            ranked_rows = sum(ranked_book[name].max_row - 4 for name in ("Weighted Sum", "TOPSIS", "VIKOR"))
            self.assertEqual(summary["total_players_scouted"], 2)
            self.assertEqual(ranked_rows, 6)
            self.assertTrue(Path(summary["rankings_workbook_file"]).exists())
            self.assertEqual(len(load_workbook(summary["rankings_workbook_file"]).sheetnames), 8)
            ranked_book.close()

    def test_all_three_methods_rank_toy_players_deterministically(self):
        players = [
            PlayerProfile("2", "B", "FWD", {"goals_per90": 0.5, "xg_per90": 0.6,
                "shots_on_target_rate": 0.5, "goals_assists_per90": 0.6, "minutes": 300}),
            PlayerProfile("1", "A", "FWD", {"goals_per90": 0.8, "xg_per90": 0.7,
                "shots_on_target_rate": 0.7, "goals_assists_per90": 0.9, "minutes": 400}),
            PlayerProfile("3", "C", "FWD", {"goals_per90": 0.2, "xg_per90": 0.3,
                "shots_on_target_rate": 0.2, "goals_assists_per90": 0.3, "minutes": 200}),
        ]
        first = _rank_group(players)
        second = _rank_group(players)
        self.assertEqual(tuple(first), ALGORITHMS)
        self.assertEqual(first, second)
        for ranked in first.values():
            self.assertEqual([entry["rank"] for entry in ranked], [1, 2, 3])
            self.assertEqual(ranked[0]["player_id"], "1")

    def test_statsbomb_dataset_has_real_role_groups_and_valid_metrics(self):
        players = load_player_pool(DEFAULT_DATASET)
        self.assertGreater(len(players), 250)
        self.assertEqual({player.position for player in players}, {"GK", "DEF", "MID", "FWD"})
        self.assertTrue(all(player.metrics["minutes"] >= 180 for player in players))
        self.assertTrue(all(0 <= player.metrics["pass_completion"] <= 1 for player in players))

    def test_real_data_runs_all_methods_and_preserves_players(self):
        players = load_player_pool(DEFAULT_DATASET)
        compute = ElasticComputePool(min_workers=1, max_workers=4)
        try:
            rankings = run_hpc_processing(players, compute)
        finally:
            compute.shutdown()
        self.assertEqual(set(rankings), {"GK", "DEF", "MID", "FWD"})
        for position, methods in rankings.items():
            self.assertEqual(set(methods), set(ALGORITHMS))
            for ranked in methods.values():
                self.assertEqual(len(ranked), sum(player.position == position for player in players))
                self.assertEqual([entry["rank"] for entry in ranked], list(range(1, len(ranked) + 1)))


if __name__ == "__main__":
    unittest.main()
